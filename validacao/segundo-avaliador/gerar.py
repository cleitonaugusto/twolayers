"""Monta o pacote do SEGUNDO AVALIADOR (avaliação às cegas dos PDFs acusados de texto oculto na rodada 5).

    PYTHONPATH=~/kit-auditoria-mcp/src python3 -I validacao/segundo-avaliador/gerar.py <resultado-da-rodada.json> <pasta-de-saida>

Gera, na pasta de saída (fora do git; os PDFs são públicos, do Govdocs1, mas grandes):
  pdfs/<código>.pdf      os arquivos acusados, renomeados por código (o nome original não aparece)
  avaliacao.xlsx         uma linha por arquivo: página e trecho para procurar; o avaliador responde se o trecho aparece
E, aqui no repositório, gabarito-interno.json: o que a ferramenta concluiu de cada código (NÃO vai para o avaliador).

O avaliador não vê o motivo da acusação nem a classificação da ferramenta: só o trecho a procurar.
"""

import json
import os
import random
import re
import shutil
import sys
from pathlib import Path

RAIZ_GOVDOCS = Path(os.environ.get("AUDITAR_GOVDOCS", Path(__file__).resolve().parents[1] / "govdocs1" / "arquivos"))
AQUI = Path(__file__).resolve().parent
EVIDENCIA = re.compile(r'^pág\. (\d+) — (.+?): "(.*)"$', re.S)
OCULTO_RECONHECIDO = re.compile(r"VerDate|\.indd|QuickTime|Template Update|recommendations made herein", re.I)


def trecho_para_busca(texto: str) -> str:
    """Pedaço curto e "buscável" com Ctrl+F: sem quebras de linha nem glifos sem texto."""
    limpo = " ".join(re.sub(r"\(cid:\d+\)", " ", texto).split())
    palavras = limpo.split(" ")
    return " ".join(palavras[:8])[:60]


def classificacao_interna(arquivo: Path, evidencias: list[str]) -> str:
    """Repete a decisão da ferramenta para dizer POR QUE o arquivo ficou acusado (só para a comparação final)."""
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src"))
    from twolayers import peca as kp, render

    if any(OCULTO_RECONHECIDO.search(e) for e in evidencias):
        return "oculto_reconhecido"
    vereditos = []
    original = kp._conferir_na_pagina

    def espiao(caminho, n, trechos, rot=0, mb=(0, 0, 0, 0)):
        alvos = [t for t in trechos if t.motivo in kp.CONFERIVEIS_NA_PAGINA and t.caixas and kp.substancial(t.texto)]
        if alvos:
            v = render.visibilidade(caminho, n, [kp._para_coordenadas_do_pdf(c, rot, mb) for t in alvos for c in t.caixas])
            i = 0
            for t in alvos:
                parte = v[i:i + len(t.caixas)]
                i += len(t.caixas)
                vereditos.append("nao_avaliou" if all(x is None for x in parte) else "invisivel")
        original(caminho, n, trechos, rot, mb)

    kp._conferir_na_pagina = espiao
    try:
        kp.extrair_pdf(arquivo)
    finally:
        kp._conferir_na_pagina = original
    if not vereditos:
        return "oculto_estrutural"
    if "nao_avaliou" in vereditos:
        return "pagina_nao_avaliou" if "invisivel" not in vereditos else "misto"
    return "pagina_confirmou_invisivel"


def main(resultado: Path, saida: Path, interno: Path = AQUI) -> int:
    import openpyxl  # noqa: F401  (falha cedo se faltar)
    rodada = json.loads(resultado.read_text(encoding="utf-8"))
    acusados = [r for r in rodada["resultados"] if r.get("achados", {}).get("T1") == "RISCO"]
    random.Random(2026).shuffle(acusados)  # ordem embaralhada e fixa (reproduzível)
    (saida / "pdfs").mkdir(parents=True, exist_ok=True)
    linhas, gabarito = [], {}
    for k, r in enumerate(acusados, 1):
        codigo = f"A{k:02d}"
        original = next(RAIZ_GOVDOCS.rglob(r["arquivo"]))
        shutil.copy(original, saida / "pdfs" / f"{codigo}.pdf")
        primeira = next((EVIDENCIA.match(e) for e in r["evidencias"]["T1"] if EVIDENCIA.match(e)), None)
        pagina, motivo, texto = (int(primeira.group(1)), primeira.group(2), primeira.group(3)) if primeira else (1, "?", "")
        linhas.append((codigo, pagina, trecho_para_busca(texto)))
        gabarito[codigo] = {"arquivo": r["arquivo"], "pagina": pagina, "motivo": motivo, "trecho": texto[:200],
                            "classificacao": classificacao_interna(original, r["evidencias"]["T1"])}
        print(f"{codigo} {gabarito[codigo]['classificacao']}", flush=True)
    planilha(saida / "avaliacao.xlsx", linhas)
    interno.mkdir(parents=True, exist_ok=True)
    (interno / "gabarito-interno.json").write_text(json.dumps(gabarito, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"pacote em {saida}: {len(linhas)} arquivos")
    return 0


def planilha(destino: Path, linhas: list[tuple]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Avaliação"
    cab = ["Código", "Página", "Trecho para procurar (Ctrl+F)", "Achou o trecho no PDF?",
           "O trecho APARECE para quem lê a página?", "Observação (opcional)"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="14532D")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for codigo, pagina, trecho in linhas:
        ws.append([codigo, pagina, trecho, None, None, None])
    achou = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    aparece = DataValidation(type="list", formula1='"Aparece,Não aparece,Aparece em parte,Não sei"', allow_blank=True)
    ws.add_data_validation(achou)
    ws.add_data_validation(aparece)
    achou.add(f"D2:D{len(linhas) + 1}")
    aparece.add(f"E2:E{len(linhas) + 1}")
    for col, larg in zip("ABCDEF", (9, 9, 52, 16, 30, 40)):
        ws.column_dimensions[col].width = larg
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 34

    guia = wb.create_sheet("Como avaliar")
    instrucoes = [
        "COMO AVALIAR (cerca de 1 a 2 horas)",
        "",
        "Cada linha é um arquivo PDF da pasta 'pdfs' (o nome do arquivo é o Código).",
        "1. Abra o PDF num leitor comum (navegador, Adobe Reader, Okular...) e vá à Página indicada.",
        "2. Procure o 'Trecho para procurar' com Ctrl+F. Marque se achou (Sim/Não).",
        "3. Olhe a página como uma pessoa leria, sem selecionar nada: o trecho APARECE na tela?",
        "   - Aparece: dá para ler normalmente.",
        "   - Não aparece: o leitor acha o trecho, mas na tela não se vê nada (branco no branco, coberto, invisível).",
        "   - Aparece em parte: só um pedaço é legível.",
        "   - Não sei: não deu para decidir (anote o motivo em Observação).",
        "4. Dica: depois de responder, aperte Ctrl+A na página. Texto escondido costuma aparecer só quando selecionado.",
        "   Responda pelo que se vê ANTES de selecionar.",
        "",
        "Não pesquise sobre os arquivos nem pergunte a ninguém o que a ferramenta concluiu: a avaliação é às cegas.",
        "Os arquivos são documentos públicos de sites do governo dos EUA (corpus Govdocs1, de pesquisa).",
    ]
    for t in instrucoes:
        guia.append([t])
    guia["A1"].font = Font(bold=True, size=13)
    guia.column_dimensions["A"].width = 110
    wb.move_sheet("Como avaliar", offset=-1)
    wb.save(destino)


if __name__ == "__main__":
    # 3º argumento opcional: pasta do gabarito interno (padrão: esta pasta, rodada do artigo)
    sys.exit(main(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]) if len(sys.argv) > 3 else AQUI))
